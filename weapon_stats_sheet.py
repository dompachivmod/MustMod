"""
Weapon Stats Spreadsheet for MustMod (Chivalry: Medieval Warfare)

Reads the UnrealScript source (the mod's classes + the base-game AOC source in the SDK)
and writes an .xlsx with, for every attack of every weapon:
  - final damage vs every class, at every body part
  - windup, release, stamina drain (ParryDrain) and turncap

Everything is read from DefaultProperties, following the `extends` chain so values
inherited from AOC base classes are picked up too. Mod-specific settings live in the
CONFIG block below; everything under it is shared between the MustMod / ZestyMod /
GiruMod copies of this script.

Damage formula (Include/MustModPawn.uci, TakeHit path):
    generic = base * DamageType[Generic]
    dmg     = (base - generic) * LocationModifiers[part] * sum(DamageType[i] * Resistances[i]) + generic
    final   = Round(dmg)
Backstab, team damage and Vanguard combo-aggression bonuses are not applied.

Usage:  python weapon_stats_sheet.py [output.xlsx]
"""
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
GAME = r"C:\Program Files (x86)\Steam\steamapps\common\chivalrymedievalwarfare"

# ===========================================================================
# CONFIG (mod-specific)
# ===========================================================================
MOD_NAME = "MustMod"
PREFIX = "MustMod"                                  # class name prefix
MOD_CLASSES = os.path.join(HERE, "Classes")
PAWN_SOURCE = "Include/MustModPawn.uci"             # where the damage formula lives

# Target classes (FamilyInfo_<Team>_<Class>). Agatha is used; Mason is checked for drift.
TARGET_CLASSES = ["Archer", "Assassin", "ManAtArms", "Vanguard", "Knight", "King"]

# Added to the parry stamina damage after the Clamp(…, 8, 25)
PARRY_DRAIN_OFFSET = 0

EXTRA_NOTES = []

def special_damage(weapon, attack, target_class, part, dmg):
    """Hook for hardcoded damage overrides in the mod's pawn code."""
    return dmg
# ===========================================================================

SOURCE_DIRS = [
    MOD_CLASSES,
    os.path.join(GAME, "Development", "Src", "AOC", "Classes"),
]

# Enum orderings (AOCPawn.uc / AOCFamilyInfo.uc / AOCDamageType.uc)
EATTACK = ["Attack_Slash", "Attack_Overhead", "Attack_Stab", "Attack_Sprint",
           "Attack_Parry", "Attack_Shove", "Attack_AltSlash", "Attack_AltOverhead"]
EHIT = ["EHIT_Head", "EHIT_Torso", "EHIT_Arm", "EHIT_Legs", "EHIT_Miss"]
EDMG = ["EDMG_Swing", "EDMG_Pierce", "EDMG_Blunt", "EDMG_Generic"]
ENUM_INDEX = {name: i for enum in (EATTACK, EHIT, EDMG) for i, name in enumerate(enum)}

# Attack slots to report (index into AttackTypeInfo / Windup / Release animations).
# Slot 4 is the parry, which never deals damage.
ATTACKS = [(0, "Slash"), (1, "Overhead"), (2, "Stab"), (3, "Sprint"), (5, "Shove")]
BODY_PARTS = [("Head", "EHIT_Head"), ("Torso", "EHIT_Torso"),
              ("Arm", "EHIT_Arm"), ("Legs", "EHIT_Legs")]

# Sprint attacks always drain this much on parry (AOCWeapon.CalculateParryDamage)
SPRINT_PARRY_DRAIN = 35


# ---------------------------------------------------------------------------
# UnrealScript DefaultProperties parsing
# ---------------------------------------------------------------------------

_file_index = None


def find_class_file(class_name):
    global _file_index
    if _file_index is None:
        _file_index = {}
        for d in reversed(SOURCE_DIRS):  # earlier dirs win
            if not os.path.isdir(d):
                print(f"WARNING: source dir not found: {d}")
                continue
            for f in os.listdir(d):
                if f.lower().endswith(".uc"):
                    _file_index[f[:-3].lower()] = os.path.join(d, f)
    return _file_index.get(class_name.lower())


def strip_comments(text):
    text = re.sub(r"/\*.*?\*/", "", text, flags=re.S)
    return re.sub(r"//[^\n]*", "", text)


def split_top_level(s, sep=","):
    parts, depth, cur, in_str = [], 0, [], False
    for ch in s:
        if ch == '"':
            in_str = not in_str
        elif not in_str and ch in "({":
            depth += 1
        elif not in_str and ch in ")}":
            depth -= 1
        if ch == sep and depth == 0 and not in_str:
            parts.append("".join(cur))
            cur = []
        else:
            cur.append(ch)
    if cur:
        parts.append("".join(cur))
    return parts


def parse_value(raw):
    raw = raw.strip().rstrip(";").strip()
    if raw.startswith("{") and raw.endswith("}"):
        raw = raw[1:-1].strip()
    if raw.startswith("(") and raw.endswith(")"):
        out = {}
        for part in split_top_level(raw[1:-1]):
            if "=" in part:
                k, v = part.split("=", 1)
                out[k.strip().lower()] = parse_value(v)
        return out
    raw = raw.strip('"')
    m = re.fullmatch(r"([-+]?\d*\.?\d+)f?", raw)
    if m:
        return float(m.group(1))
    return raw


def parse_class_file(path):
    """Return (parent_class_name, {prop_lower: value or {index: value}})."""
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        text = strip_comments(f.read())

    m = re.search(r"\bclass\s+\w+\s+extends\s+(\w+)", text, re.I)
    parent = m.group(1) if m else None

    m = re.search(r"\bdefaultproperties\b", text, re.I)
    props = {}
    if not m:
        return parent, props
    body = text[m.end():]
    body = body[body.find("{") + 1:body.rfind("}")]  # drop the block's own braces
    body = re.sub(r"^\s*begin\s+object.*?^\s*end\s+object", "", body, flags=re.I | re.S | re.M)

    # Join physical lines into statements until brackets balance
    statements, cur, depth = [], "", 0
    for line in body.splitlines():
        cur += " " + line.strip()
        depth += line.count("(") + line.count("{") - line.count(")") - line.count("}")
        if depth <= 0:
            statements.append(cur.strip())
            cur, depth = "", 0

    stmt_re = re.compile(r"^(\w+)\s*(?:[\(\[]\s*(\w+)\s*[\)\]])?\s*=\s*(.*)$")
    for st in statements:
        sm = stmt_re.match(st)
        if not sm:
            continue
        key, idx, val = sm.group(1).lower(), sm.group(2), parse_value(sm.group(3))
        if idx is None:
            props[key] = val
        else:
            idx = int(idx) if idx.isdigit() else ENUM_INDEX.get(idx, idx)
            arr = props.setdefault(key, {})
            if not isinstance(arr, dict):
                arr = props[key] = {}
            # partial struct assignment merges onto what's already there
            if isinstance(val, dict) and isinstance(arr.get(idx), dict):
                merged = dict(arr[idx]); merged.update(val); val = merged
            arr[idx] = val
    return parent, props


_resolved = {}


def resolve(class_name):
    """Merged DefaultProperties for a class, including all parents we have source for."""
    key = class_name.lower()
    if key in _resolved:
        return _resolved[key]
    path = find_class_file(class_name)
    if path is None:
        _resolved[key] = {}
        return {}
    parent, own = parse_class_file(path)
    merged = {}
    if parent:
        for k, v in resolve(parent).items():
            merged[k] = dict(v) if isinstance(v, dict) else v
    for k, v in own.items():
        if isinstance(v, dict) and isinstance(merged.get(k), dict) and all(isinstance(i, int) for i in v):
            for i, item in v.items():
                base = merged[k].get(i)
                if isinstance(item, dict) and isinstance(base, dict):
                    item = {**base, **item}
                merged[k][i] = item
        else:
            merged[k] = v
    merged["__file__"] = path
    merged["__chain__"] = [class_name.lower()] + (resolve(parent).get("__chain__", []) if parent else [])
    _resolved[key] = merged
    return merged


def class_ref(value):
    """class'Foo' / "AOC.Foo" / Foo -> Foo"""
    if not isinstance(value, str) or not value or value.lower() == "none":
        return None
    m = re.search(r"'([^']+)'", value)
    name = m.group(1) if m else value
    return name.split(".")[-1]


def arr_get(props, key, idx, default=None):
    arr = props.get(key.lower())
    if isinstance(arr, dict):
        return arr.get(idx, default)
    return default


# ---------------------------------------------------------------------------
# Game data
# ---------------------------------------------------------------------------

def load_classes():
    classes = {}
    for name in TARGET_CLASSES:
        stats = {}
        for team in ("Agatha", "Mason"):
            cls = f"{PREFIX}FamilyInfo_{team}_{name}"
            if not find_class_file(cls):
                continue
            p = resolve(cls)
            stats[team] = {
                "res": [float(arr_get(p, "DamageResistances", i, 1.0)) for i in range(4)],
                "loc": {part: float(arr_get(p, "LocationModifiers", ENUM_INDEX[e], 1.0))
                        for part, e in BODY_PARTS},
                "file": cls,
            }
        if not stats:
            print(f"WARNING: no FamilyInfo found for {name}, skipping")
            continue
        if len(stats) == 2 and (stats["Agatha"]["res"] != stats["Mason"]["res"]
                                or stats["Agatha"]["loc"] != stats["Mason"]["loc"]):
            print(f"WARNING: Agatha/Mason {name} stats differ; using Agatha")
        classes[name] = stats.get("Agatha") or stats["Mason"]
    return classes


_dmg_cache = {}


def damage_type_mix(dmg_class):
    if dmg_class not in _dmg_cache:
        p = resolve(dmg_class)
        if not p:
            print(f"WARNING: damage type {dmg_class} not found, treating as Generic")
            _dmg_cache[dmg_class] = [0.0, 0.0, 0.0, 1.0]
        else:
            _dmg_cache[dmg_class] = [float(arr_get(p, "DamageType", i, 0.0)) for i in range(4)]
    return _dmg_cache[dmg_class]


def final_damage(base, mix, cls, part):
    generic = base * mix[3]
    resist = sum(mix[i] * cls["res"][i] for i in range(4))
    return round((base - generic) * cls["loc"][part] * resist + generic)


def anim_len(props, key, idx):
    a = arr_get(props, key, idx)
    if isinstance(a, dict):
        v = a.get("fanimationlength")
        return float(v) if isinstance(v, float) else None
    return None


def turncaps(props, attack_idx):
    """Horizontal/vertical rotate speed while performing the attack (AOCWeapon.GetHorizontalRotateSpeed)."""
    def f(k):
        v = props.get(k.lower())
        return float(v) if isinstance(v, float) else None

    if attack_idx == 3:
        return f("SprintAttackHorizRotateSpeed"), f("SprintAttackVerticalRotateSpeed")
    horiz, vert = f("AttackHorizRotateSpeed"), f("VerticalRotateSpeed")
    # Per-attack overrides (e.g. MustModWeapon_GrandLance)
    prefix = {0: "Slash", 1: "Overhead", 2: "Stab"}.get(attack_idx)
    if prefix:
        h, v = f(prefix + "HorizRotateSpeed"), f(prefix + "VertRotateSpeed")
        horiz = h if h else horiz
        vert = v if v else vert
    return horiz, vert


def load_weapons():
    weapons = []
    for path in sorted(os.listdir(MOD_CLASSES)):
        m = re.fullmatch(PREFIX + r"Weapon_(\w+)\.uc", path, re.I)
        if not m:
            continue
        wclass = path[:-3]
        wp = resolve(wclass)
        if "aocrangeweapon" in wp.get("__chain__", []):
            continue  # bows, crossbows, throwables: damage lives in the projectile class
        att_class = class_ref(wp.get("attachmentclass"))
        ap = resolve(att_class) if att_class else {}
        if not isinstance(ap.get("attacktypeinfo"), dict):
            continue  # shields, bows etc. with no melee attack table

        attacks = []
        for idx, label in ATTACKS:
            info = arr_get(ap, "AttackTypeInfo", idx)
            if not isinstance(info, dict):
                continue
            base = float(info.get("fbasedamage", 0.0) or 0.0)
            if base <= 0:
                continue
            dmg_class = class_ref(info.get("cdamagetype")) or "AOCDmgType_Generic"
            if idx == 3:
                drain = SPRINT_PARRY_DRAIN + PARRY_DRAIN_OFFSET
            else:
                d = arr_get(wp, "ParryDrain", idx)
                drain = float(d) if isinstance(d, float) else None
            h, v = turncaps(wp, idx)
            attacks.append({
                "idx": idx, "name": label, "base": base, "dmg_class": dmg_class,
                "windup": anim_len(wp, "WindupAnimations", idx),
                "release": anim_len(wp, "ReleaseAnimations", idx),
                "recovery": anim_len(wp, "RecoveryAnimations", idx),
                "drain": drain, "turn_h": h, "turn_v": v,
            })
        if attacks:
            weapons.append({
                "name": m.group(1), "class": wclass, "attachment": att_class,
                "parry_negation": wp.get("fparrynegation"), "attacks": attacks,
            })
    return weapons


# ---------------------------------------------------------------------------
# Spreadsheet
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
SUB_FILL = PatternFill("solid", fgColor="374151")
BAND_FILL = PatternFill("solid", fgColor="F3F4F6")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="9CA3AF")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_workbook(weapons, classes, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Weapon Stats"

    info_cols = ["Weapon", "Attack", "Damage Type", "Base Dmg",
                 "Windup (s)", "Release (s)", "Recovery (s)",
                 "Stam Drain (ParryDrain)", "Turncap Horiz", "Turncap Vert"]
    n_info = len(info_cols)

    # Row 1: group headers, Row 2: column headers
    for c, title in enumerate(info_cols, 1):
        ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
        ws.cell(1, c, title)
    col = n_info + 1
    for cname in classes:
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(BODY_PARTS) - 1)
        ws.cell(1, col, cname)
        for j, (part, _) in enumerate(BODY_PARTS):
            ws.cell(2, col + j, part)
        col += len(BODY_PARTS)
    last_col = col - 1

    for r in (1, 2):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.fill = HEADER_FILL if r == 1 else SUB_FILL
            cell.font = WHITE_BOLD
            cell.alignment = CENTER

    row = 3
    for wi, w in enumerate(weapons):
        band = BAND_FILL if wi % 2 == 0 else None
        for ai, a in enumerate(w["attacks"]):
            mix = damage_type_mix(a["dmg_class"])
            values = [w["name"], a["name"], a["dmg_class"].replace("AOCDmgType_", ""), a["base"],
                      a["windup"], a["release"], a["recovery"],
                      a["drain"], a["turn_h"], a["turn_v"]]
            for cname, cls in classes.items():
                for part, _ in BODY_PARTS:
                    dmg = final_damage(a["base"], mix, cls, part)
                    values.append(special_damage(w["name"], a["name"], cname, part, dmg))
            for c, v in enumerate(values, 1):
                cell = ws.cell(row, c, v)
                if band:
                    cell.fill = band
                if c > 3:
                    cell.alignment = Alignment(horizontal="center")
                if ai == 0:
                    cell.border = Border(top=THIN)
            row += 1

    # Vertical separators between class groups
    for c in range(n_info + 1, last_col + 1, len(BODY_PARTS)):
        for r in range(1, row):
            cell = ws.cell(r, c)
            cell.border = Border(left=THIN, top=cell.border.top)

    widths = [16, 10, 13, 9, 9, 9, 10, 12, 10, 10]
    for c, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = wdt
    for c in range(n_info + 1, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 7
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = ws.cell(3, 3)
    ws.auto_filter.ref = f"A2:{get_column_letter(last_col)}{row - 1}"

    # Reference sheet: the inputs the numbers came from
    ref = wb.create_sheet("Reference")
    ref.append(["Class", "Source", "Swing res", "Pierce res", "Blunt res", "Generic res"]
               + [f"{p} mod" for p, _ in BODY_PARTS])
    for cname, cls in classes.items():
        ref.append([cname, cls["file"]] + cls["res"] + [cls["loc"][p] for p, _ in BODY_PARTS])
    ref.append([])
    ref.append(["Damage Type", "Swing", "Pierce", "Blunt", "Generic"])
    for dname, mix in sorted(_dmg_cache.items()):
        ref.append([dname] + mix)
    ref.append([])
    ref.append(["Weapon", "Weapon class", "Attachment class", "fParryNegation (as defender)"])
    for w in weapons:
        ref.append([w["name"], w["class"], w["attachment"], w["parry_negation"]])
    ref.append([])
    for note in [
        f"Damage = Round((Base - Generic) * BodyPartMod * sum(DmgTypeMix * ClassRes) + Generic), from {PAWN_SOURCE}.",
        "Not applied: backstab modifier, team damage, Vanguard combo aggression bonus.",
        "Stam Drain is the attacker's raw ParryDrain. Actual drain on the parrier = Clamp(Round(ParryDrain - defender's fParryNegation), 8, 25)"
        + (f" {PARRY_DRAIN_OFFSET:+d}" if PARRY_DRAIN_OFFSET else "")
        + f". Sprint attacks always drain {SPRINT_PARRY_DRAIN + PARRY_DRAIN_OFFSET} (shown in the Sprint row).",
        "Turncap = AttackHorizRotateSpeed / VerticalRotateSpeed while attacking (Sprint uses the SprintAttack* values). Per-attack overrides (GrandLance) are applied.",
        "Windup/Release/Recovery = fAnimationLength of WindupAnimations / ReleaseAnimations / RecoveryAnimations for that attack slot.",
    ] + EXTRA_NOTES:
        ref.append([note])
    for c in range(1, ref.max_column + 1):
        ref.cell(1, c).font = Font(bold=True)
    ref.column_dimensions["A"].width = 22
    ref.column_dimensions["B"].width = 32
    ref.column_dimensions["C"].width = 34

    wb.save(out_path)


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "weapon_stats.xlsx")
    classes = load_classes()
    weapons = load_weapons()
    write_workbook(weapons, classes, out)
    n_attacks = sum(len(w["attacks"]) for w in weapons)
    print(f"{MOD_NAME}: wrote {len(weapons)} weapons, {n_attacks} attacks, {len(classes)} classes -> {out}")


if __name__ == "__main__":
    main()
